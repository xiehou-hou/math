# -*- coding: utf-8 -*-
"""提取附件1/附件2的 docx 文本内容"""
import os
import docx

FILES = [
    (r"d:\数学建模\B题\附件\附件1.docx", r"d:\数学建模\B题\附件1_文本.txt"),
    (r"d:\数学建模\B题\附件\附件2.docx", r"d:\数学建模\B题\附件2_文本.txt"),
]


def extract(src: str, dst: str) -> None:
    if not os.path.exists(src):
        print(f"文件不存在: {src}")
        return
    doc = docx.Document(src)
    with open(dst, "w", encoding="utf-8") as fout:
        for i, para in enumerate(doc.paragraphs, 1):
            fout.write(para.text + "\n")
        # 表格内容
        for ti, table in enumerate(doc.tables, 1):
            fout.write(f"\n----- 表格 {ti} -----\n")
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                fout.write(" | ".join(cells) + "\n")
    size = os.path.getsize(dst)
    print(f"{src} -> {dst}  大小: {size} 字节")


def main():
    for src, dst in FILES:
        extract(src, dst)


if __name__ == "__main__":
    main()
